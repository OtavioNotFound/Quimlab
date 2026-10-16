from docx import Document

path = r"C:\Users\Otavi\Downloads\Briefing Da Empresa Prato Certo.docx"
doc = Document(path)

print("PARAGRAFOS")
for index, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.strip():
        print(f"{index}: {paragraph.text}")

print("TABELAS")
for table_index, table in enumerate(doc.tables, start=1):
    print(f"TABELA {table_index}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
